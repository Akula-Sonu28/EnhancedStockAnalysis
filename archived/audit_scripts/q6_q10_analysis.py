#!/usr/bin/env python3
"""Q6–Q10 deep analysis of the most recent Enhanced Stock Report."""

import glob, os, math
import openpyxl
from collections import Counter, defaultdict

REPORTS_DIR = os.path.join(os.path.dirname(__file__), "reports")

def load_latest_report():
    files = sorted(glob.glob(os.path.join(REPORTS_DIR, "*.xlsx")),
                   key=os.path.getmtime, reverse=True)
    if not files:
        raise FileNotFoundError("No .xlsx in reports/")
    path = files[0]
    print(f"=== Report: {os.path.basename(path)} ===\n")
    return openpyxl.load_workbook(path, read_only=True, data_only=True), path


def sheet_to_dicts(ws, header_row=1):
    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    headers = rows[0]
    return [dict(zip(headers, r)) for r in rows[1:]]


def safe_float(v):
    if v is None:
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


def banner(title):
    print("\n" + "=" * 72)
    print(f"  {title}")
    print("=" * 72)


# ── helpers ──────────────────────────────────────────────────────────────

def pearson(xs, ys):
    pairs = [(x, y) for x, y in zip(xs, ys)
             if x is not None and y is not None]
    if len(pairs) < 3:
        return None, len(pairs)
    n = len(pairs)
    sx = sum(x for x, _ in pairs)
    sy = sum(y for _, y in pairs)
    sxx = sum(x * x for x, _ in pairs)
    syy = sum(y * y for _, y in pairs)
    sxy = sum(x * y for x, y in pairs)
    denom = math.sqrt((n * sxx - sx ** 2) * (n * syy - sy ** 2))
    if denom == 0:
        return 0.0, n
    return (n * sxy - sx * sy) / denom, n


def text_histogram(values, bucket_size=5, width=50):
    if not values:
        print("  (no data)")
        return
    lo = int(min(values) // bucket_size) * bucket_size
    hi = int(max(values) // bucket_size + 1) * bucket_size
    buckets = defaultdict(int)
    for v in values:
        b = int(v // bucket_size) * bucket_size
        buckets[b] += 1
    max_count = max(buckets.values()) if buckets else 1
    for b in range(lo, hi + 1, bucket_size):
        c = buckets[b]
        bar = "█" * max(1, int(c / max_count * width)) if c else ""
        label = f"  {b:5.0f}–{b + bucket_size - 1:5.0f}"
        print(f"{label} | {bar} {c}")


def percentile(sorted_vals, p):
    if not sorted_vals:
        return None
    k = (len(sorted_vals) - 1) * p / 100
    f = int(k)
    c = f + 1 if f + 1 < len(sorted_vals) else f
    d = k - f
    return sorted_vals[f] + d * (sorted_vals[c] - sorted_vals[f])


# ═════════════════════════════════════════════════════════════════════════
#  MAIN
# ═════════════════════════════════════════════════════════════════════════

wb, fpath = load_latest_report()

# ── load Complete Data ───────────────────────────────────────────────────
cd = sheet_to_dicts(wb["Complete Data"])
print(f"Complete Data: {len(cd)} stocks")

# ── load Portfolio Allocation (header in row 2) ─────────────────────────
pa_ws = wb["Portfolio Allocation"]
pa_rows = list(pa_ws.iter_rows(values_only=True))
pa_headers = list(pa_rows[1])  # row 2 = actual headers
pa_data = [dict(zip(pa_headers, r)) for r in pa_rows[2:] if any(r)]
print(f"Portfolio Allocation: {len(pa_data)} rows")


# ═════════════════════════════════════════════════════════════════════════
#  Q6: MTF value or noise?
# ═════════════════════════════════════════════════════════════════════════
banner("Q6: Is MTF (multi-timeframe) adding value or noise?")

mtf_cols = [h for h in cd[0].keys() if h and ("mtf" in h.lower() or "multi_timeframe" in h.lower())]
print(f"\nMTF-related columns ({len(mtf_cols)}):")
for c in mtf_cols:
    print(f"  • {c}")

# MTF status distribution
mtf_status = Counter(str(r.get("mtf_analysis_status", "MISSING")) for r in cd)
print(f"\nmtf_analysis_status distribution:")
for k, v in mtf_status.most_common():
    print(f"  {k:20s} : {v:4d}  ({v/len(cd)*100:.1f}%)")

# Score clustering for MTF success vs failure
success_scores = [safe_float(r.get("overall_score")) for r in cd
                  if str(r.get("mtf_analysis_status")).lower() == "success"]
fail_scores = [safe_float(r.get("overall_score")) for r in cd
               if str(r.get("mtf_analysis_status")).lower() != "success"]
success_scores = [s for s in success_scores if s is not None]
fail_scores = [s for s in fail_scores if s is not None]

if success_scores:
    print(f"\nOverall scores where MTF succeeded ({len(success_scores)} stocks):")
    print(f"  mean={sum(success_scores)/len(success_scores):.2f}  "
          f"median={sorted(success_scores)[len(success_scores)//2]:.2f}  "
          f"min={min(success_scores):.2f}  max={max(success_scores):.2f}")
if fail_scores:
    print(f"\nOverall scores where MTF FAILED ({len(fail_scores)} stocks):")
    print(f"  mean={sum(fail_scores)/len(fail_scores):.2f}  "
          f"median={sorted(fail_scores)[len(fail_scores)//2]:.2f}  "
          f"min={min(fail_scores):.2f}  max={max(fail_scores):.2f}")

# Correlation: mtf_composite_score vs overall_score
mtf_scores = [safe_float(r.get("mtf_composite_score")) for r in cd]
overall_scores = [safe_float(r.get("overall_score")) for r in cd]
r_val, n_pairs = pearson(mtf_scores, overall_scores)
if r_val is not None:
    print(f"\nCorrelation (mtf_composite_score ↔ overall_score): r = {r_val:.4f}  (n={n_pairs})")
    if abs(r_val) < 0.2:
        print("  → WEAK correlation — MTF may not be contributing meaningful signal")
    elif abs(r_val) < 0.5:
        print("  → MODERATE correlation — MTF adds some signal")
    else:
        print("  → STRONG correlation — MTF is a key driver")

# Also check mtf_composite_score_final
mtf_final = [safe_float(r.get("mtf_composite_score_final")) for r in cd]
r2, n2 = pearson(mtf_final, overall_scores)
if r2 is not None:
    print(f"Correlation (mtf_composite_score_final ↔ overall_score): r = {r2:.4f}  (n={n2})")


# ═════════════════════════════════════════════════════════════════════════
#  Q7: Stale holdings (score=0) unfairly treated?
# ═════════════════════════════════════════════════════════════════════════
banner("Q7: Are stale holdings (score=0) being unfairly treated?")

# Check in Complete Data
zero_risk = [r for r in cd if safe_float(r.get("risk_adjusted_score")) == 0]
zero_overall = [r for r in cd if safe_float(r.get("overall_score")) == 0]
print(f"\nIn Complete Data:")
print(f"  Stocks with risk_adjusted_score = 0 : {len(zero_risk)}")
print(f"  Stocks with overall_score = 0       : {len(zero_overall)}")

all_zero_symbols = set()
for r in zero_risk + zero_overall:
    all_zero_symbols.add(r.get("symbol"))

if all_zero_symbols:
    print(f"\n  Zero-score stocks ({len(all_zero_symbols)}):")
    for sym in sorted(all_zero_symbols):
        row = next((r for r in cd if r.get("symbol") == sym), {})
        ras = row.get("risk_adjusted_score")
        os_ = row.get("overall_score")
        rec = row.get("final_recommendation", row.get("corrected_recommendation", "N/A"))
        status = row.get("status", "?")
        fund = row.get("fundamental_status", "?")
        tech = row.get("real_technical_status", "?")
        print(f"    {sym:20s}  risk_adj={ras}  overall={os_}  rec={rec}  "
              f"status={status}  fund={fund}  tech={tech}")
else:
    print("  → No zero-score stocks found in Complete Data.")

# Check Portfolio Allocation for zero scores
print(f"\nIn Portfolio Allocation sheet:")
pa_zero = [r for r in pa_data
           if safe_float(r.get("RISK SC")) == 0 or safe_float(r.get("SCORE")) == 0]
if pa_zero:
    print(f"  Stocks with SCORE=0 or RISK SC=0 ({len(pa_zero)}):")
    for r in pa_zero:
        print(f"    {str(r.get('symbol','?')):20s}  SCORE={r.get('SCORE')}  "
              f"ADJ_SCORE={r.get('ADJ SCORE')}  RISK_SC={r.get('RISK SC')}  "
              f"ACTION={r.get('ACTION')}  RISK={r.get('RISK')}")
else:
    print("  → No zero-score stocks in Portfolio Allocation either.")

# Also check very low scores (< 10) — these might be quasi-stale
very_low = [r for r in cd if (safe_float(r.get("overall_score")) or 999) < 10]
if very_low:
    print(f"\n  Stocks with overall_score < 10 ({len(very_low)}) — possibly stale/broken:")
    for r in very_low:
        print(f"    {str(r.get('symbol','')):20s}  overall={r.get('overall_score')}  "
              f"rec={r.get('final_recommendation','?')}  status={r.get('status','?')}")


# ═════════════════════════════════════════════════════════════════════════
#  Q8: Crisis detector adding noise in non-crisis periods?
# ═════════════════════════════════════════════════════════════════════════
banner("Q8: Is the crisis detector adding noise in non-crisis periods?")

crisis_cols = [h for h in cd[0].keys() if h and "crisis" in h.lower()]
print(f"\nCrisis-related columns ({len(crisis_cols)}):")
for c in crisis_cols:
    print(f"  • {c}")

crisis_types = Counter(str(r.get("crisis_type_detected", "MISSING")) for r in cd)
print(f"\ncrisis_type_detected distribution:")
for k, v in crisis_types.most_common():
    print(f"  {k:30s} : {v:4d}  ({v/len(cd)*100:.1f}%)")

severity_dist = Counter(str(r.get("crisis_severity", "MISSING")) for r in cd)
print(f"\ncrisis_severity distribution:")
for k, v in severity_dist.most_common():
    print(f"  {k:30s} : {v:4d}  ({v/len(cd)*100:.1f}%)")

crisis_adj = [safe_float(r.get("crisis_score_adjustment")) for r in cd]
nonzero_adj = [v for v in crisis_adj if v is not None and v != 0]
print(f"\ncrisis_score_adjustment:")
print(f"  Total stocks         : {len(cd)}")
print(f"  Non-zero adjustments : {len(nonzero_adj)}")
if nonzero_adj:
    print(f"  Range                : {min(nonzero_adj):.2f} to {max(nonzero_adj):.2f}")
    print(f"  Mean adjustment      : {sum(nonzero_adj)/len(nonzero_adj):.2f}")

# KEY CHECK: Is crisis NONE but still adjusting?
none_with_adj = [r for r in cd
                 if str(r.get("crisis_type_detected", "")).upper() in ("NONE", "NORMAL", "NO_CRISIS")
                 and safe_float(r.get("crisis_score_adjustment")) not in (None, 0)]
print(f"\n⚠  Stocks with crisis='NONE/NORMAL' BUT non-zero adjustment: {len(none_with_adj)}")
if none_with_adj:
    print("  THIS IS A BUG — adjustments should be 0 when no crisis detected!")
    for r in none_with_adj[:10]:
        print(f"    {str(r.get('symbol','')):20s}  crisis={r.get('crisis_type_detected')}  "
              f"severity={r.get('crisis_severity')}  adj={r.get('crisis_score_adjustment')}")
    if len(none_with_adj) > 10:
        print(f"    ... and {len(none_with_adj) - 10} more")
else:
    print("  ✓ No phantom adjustments when crisis=NONE — good.")


# ═════════════════════════════════════════════════════════════════════════
#  Q9: BUY/INCREASE recommendations for HIGH/VERY HIGH risk stocks?
# ═════════════════════════════════════════════════════════════════════════
banner("Q9: How many stocks are recommended BUY but have HIGH/VERY HIGH risk?")

risk_col = "risk_category"
rec_col = "final_recommendation"

risk_dist = Counter(str(r.get(risk_col, "MISSING")) for r in cd)
print(f"\nRisk category distribution:")
for k, v in risk_dist.most_common():
    print(f"  {k:20s} : {v:4d}  ({v/len(cd)*100:.1f}%)")

rec_dist = Counter(str(r.get(rec_col, "MISSING")) for r in cd)
print(f"\nFinal recommendation distribution:")
for k, v in rec_dist.most_common():
    print(f"  {k:20s} : {v:4d}  ({v/len(cd)*100:.1f}%)")

high_risk_labels = {"HIGH", "VERY HIGH", "VERY_HIGH", "EXTREME"}
buy_labels = {"BUY", "STRONG BUY", "STRONG_BUY", "INCREASE", "ACCUMULATE"}

risky_buys = []
for r in cd:
    risk = str(r.get(risk_col, "")).upper().strip()
    rec = str(r.get(rec_col, "")).upper().strip()
    if risk in high_risk_labels and rec in buy_labels:
        risky_buys.append(r)

print(f"\n⚠  BUY/INCREASE recommendations for HIGH+ risk stocks: {len(risky_buys)}")
if risky_buys:
    print("  These are potential TRAPS — high score but dangerous:\n")
    print(f"  {'SYMBOL':20s} {'RISK':12s} {'REC':16s} {'SCORE':>8s} {'RISK_ADJ':>10s} {'VOLATILITY':>10s}")
    print(f"  {'-'*20} {'-'*12} {'-'*16} {'-'*8} {'-'*10} {'-'*10}")
    for r in sorted(risky_buys, key=lambda x: safe_float(x.get("overall_score")) or 0, reverse=True):
        print(f"  {str(r.get('symbol','')):20s} "
              f"{str(r.get(risk_col,'')):12s} "
              f"{str(r.get(rec_col,'')):16s} "
              f"{str(r.get('overall_score','')):>8s} "
              f"{str(r.get('risk_adjusted_score','')):>10s} "
              f"{str(r.get('volatility','')):>10s}")
else:
    print("  ✓ No high-risk stocks being recommended for buying.")

# Also check the inverse: SELL on LOW risk
low_risk_sells = [r for r in cd
                  if str(r.get(risk_col, "")).upper().strip() in ("LOW", "VERY LOW", "VERY_LOW")
                  and str(r.get(rec_col, "")).upper().strip() in ("SELL", "STRONG SELL", "STRONG_SELL", "REDUCE")]
if low_risk_sells:
    print(f"\n  (Also: {len(low_risk_sells)} SELL recommendations on LOW-risk stocks — could be overcautious)")


# ═════════════════════════════════════════════════════════════════════════
#  Q10: Score distribution — too compressed or too spread?
# ═════════════════════════════════════════════════════════════════════════
banner("Q10: What's the actual score distribution?")

scores_raw = [safe_float(r.get("overall_score")) for r in cd]
scores = sorted([s for s in scores_raw if s is not None])

if not scores:
    print("  No valid overall_score data!")
else:
    n = len(scores)
    mean = sum(scores) / n
    median = scores[n // 2]
    std = math.sqrt(sum((s - mean) ** 2 for s in scores) / n)
    p25 = percentile(scores, 25)
    p75 = percentile(scores, 75)

    print(f"\n  Metric           Value")
    print(f"  ────────────────────────")
    print(f"  Count          : {n}")
    print(f"  Min            : {min(scores):.2f}")
    print(f"  25th pctl      : {p25:.2f}")
    print(f"  Median         : {median:.2f}")
    print(f"  Mean           : {mean:.2f}")
    print(f"  75th pctl      : {p75:.2f}")
    print(f"  Max            : {max(scores):.2f}")
    print(f"  Std Dev        : {std:.2f}")
    print(f"  IQR            : {p75 - p25:.2f}")

    print(f"\n  Distribution (buckets of 5 points):")
    text_histogram(scores, bucket_size=5)

    print(f"\n  Diagnosis:")
    if std < 5:
        print(f"  ⚠  STD = {std:.2f} < 5 → Scores are TOO COMPRESSED")
        print(f"     Everything looks the same — the system can't differentiate stocks.")
    elif std > 15:
        print(f"  ⚠  STD = {std:.2f} > 15 → Scores are TOO SPREAD")
        print(f"     Adjustments are dominating the base score — unstable rankings.")
    else:
        print(f"  ✓  STD = {std:.2f} is in the healthy range (5–15)")
        print(f"     The system has reasonable differentiation.")

    iqr = p75 - p25
    if iqr < 8:
        print(f"  ⚠  IQR = {iqr:.2f} < 8 → Middle 50% of stocks are very close together")
    elif iqr > 20:
        print(f"  ⚠  IQR = {iqr:.2f} > 20 → Wide spread in middle 50% — good differentiation")
    else:
        print(f"  ✓  IQR = {iqr:.2f} — reasonable spread")

    # Score bands
    bands = {"0-20": 0, "20-40": 0, "40-60": 0, "60-80": 0, "80-100": 0}
    for s in scores:
        if s < 20: bands["0-20"] += 1
        elif s < 40: bands["20-40"] += 1
        elif s < 60: bands["40-60"] += 1
        elif s < 80: bands["60-80"] += 1
        else: bands["80-100"] += 1
    print(f"\n  Score Bands:")
    for band, count in bands.items():
        pct = count / n * 100
        bar = "█" * int(pct / 2)
        print(f"    {band:8s} : {count:4d} ({pct:5.1f}%)  {bar}")

wb.close()

print("\n" + "=" * 72)
print("  Analysis complete.")
print("=" * 72)
