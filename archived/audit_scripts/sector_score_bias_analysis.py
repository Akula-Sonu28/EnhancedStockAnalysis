#!/usr/bin/env python3
"""
Sector Score Bias Analysis
Traces how individual scoring components vary by sector to identify systematic biases.
"""

import os
import glob
import openpyxl
import statistics
from collections import defaultdict

REPORTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "reports")

SCORE_COLUMNS = {
    "fundamental_score": "Fund_Raw",
    "fundamental_score_final": "Fund_Final",
    "FundamentalScore": "Fund_Display",
    "pe_relative_score": "PE_Rel",
    "pb_relative_score": "PB_Rel",
    "roe_relative_score": "ROE_Rel",
    "growth_relative_score": "Growth_Rel",
    "enhanced_technical_score_final": "EnhTech",
    "real_technical_score_final": "RealTech",
    "TechnicalScore": "Tech_Display",
    "combined_technical_score_final": "CombTech",
    "mtf_composite_score_final": "MTF",
    "institutional_score_final": "Instit",
    "volume_composite_score": "Vol_Comp",
    "hybrid_fundamental_quality": "H_Fund",
    "hybrid_momentum_technical": "H_Momentum",
    "hybrid_volume_strength": "H_Volume",
    "hybrid_multi_timeframe": "H_MTF",
    "hybrid_ml_signal": "H_ML",
    "hybrid_risk_adjustment": "H_Risk",
    "hybrid_sector_multiplier": "H_SectorMult",
    "hybrid_overall_score": "H_Overall",
    "portfolio_context_score": "Portfolio_Ctx",
    "regime_adjustment_amount": "Regime_Adj",
    "sentiment_adjustment_amount": "Sentiment_Adj",
    "volume_adjustment_amount": "Volume_Adj",
    "ml_score_adjustment": "ML_Adj",
    "sentiment_score_contribution": "Sent_Contrib",
    "volume_score_contribution": "Vol_Contrib",
    "pattern_score_contribution": "Pattern_Contrib",
    "crisis_score_adjustment": "Crisis_Adj",
    "sector_performance_adj": "Sector_PerfAdj",
    "phase1_blended_score": "Phase1_Blend",
    "final_blended_score": "Final_Blend",
    "risk_adjusted_score": "Risk_Adj_Score",
    "overall_score": "Overall",
    "optimized_score": "Optimized",
    "corrected_overall_score": "Corrected",
    "improved_overall_score": "Improved",
}

VALUATION_COLS = {
    "pe_ratio": "PE",
    "pb_ratio": "PB",
    "roe": "ROE",
    "revenue_growth": "RevGrowth",
    "debt_to_equity": "D/E",
}


def find_latest_report():
    files = glob.glob(os.path.join(REPORTS_DIR, "*.xlsx"))
    if not files:
        raise FileNotFoundError("No .xlsx reports found in reports/")
    return max(files, key=os.path.getmtime)


def load_data(filepath):
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    if "Complete Data" not in wb.sheetnames:
        raise ValueError("'Complete Data' sheet not found")
    ws = wb["Complete Data"]

    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    headers = [str(h) if h else "" for h in rows[0]]
    col_idx = {h: i for i, h in enumerate(headers)}

    data = []
    for row in rows[1:]:
        symbol = row[col_idx.get("symbol", 0)]
        sector = row[col_idx.get("sector", 6)]
        if not symbol or not sector:
            continue

        record = {"symbol": str(symbol), "sector": str(sector)}
        for col_name in list(SCORE_COLUMNS.keys()) + list(VALUATION_COLS.keys()):
            if col_name in col_idx:
                val = row[col_idx[col_name]]
                try:
                    record[col_name] = float(val) if val is not None else None
                except (ValueError, TypeError):
                    record[col_name] = None
        data.append(record)
    return data


def sector_stats(data, columns, labels):
    sectors = defaultdict(list)
    for rec in data:
        sectors[rec["sector"]].append(rec)

    results = {}
    for sector, stocks in sorted(sectors.items()):
        stats = {"count": len(stocks)}
        for col in columns:
            vals = [s[col] for s in stocks if s.get(col) is not None]
            stats[col] = statistics.mean(vals) if vals else None
        results[sector] = stats
    return results


def print_table(title, results, columns, labels, fmt=".1f"):
    print(f"\n{'=' * 120}")
    print(f"  {title}")
    print(f"{'=' * 120}")

    header = f"{'Sector':<30} {'N':>4}"
    for label in labels:
        header += f" {label:>10}"
    print(header)
    print("-" * len(header))

    all_vals = {col: [] for col in columns}

    for sector, stats in sorted(results.items(), key=lambda x: -(x[1].get(columns[-1]) or 0)):
        line = f"{sector:<30} {stats['count']:>4}"
        for col in columns:
            v = stats.get(col)
            if v is not None:
                line += f" {v:>10.{fmt[-2]}f}"
                all_vals[col].append(v)
            else:
                line += f" {'N/A':>10}"
        print(line)

    print("-" * len(header))
    line = f"{'OVERALL MEAN':<30} {'':>4}"
    for col in columns:
        vals = all_vals[col]
        if vals:
            line += f" {statistics.mean(vals):>10.{fmt[-2]}f}"
        else:
            line += f" {'N/A':>10}"
    print(line)

    line = f"{'SPREAD (max-min)':<30} {'':>4}"
    for col in columns:
        vals = all_vals[col]
        if len(vals) >= 2:
            spread = max(vals) - min(vals)
            line += f" {spread:>10.{fmt[-2]}f}"
        else:
            line += f" {'N/A':>10}"
    print(line)


def gap_analysis(results, columns, labels):
    print(f"\n{'=' * 120}")
    print("  COMPONENT GAP ANALYSIS: Which component creates the biggest sector disparity?")
    print(f"{'=' * 120}")
    print(f"{'Component':<25} {'Min Sector':<25} {'Min Val':>8} {'Max Sector':<25} {'Max Val':>8} {'Spread':>8} {'CV%':>8}")
    print("-" * 107)

    gaps = []
    for col, label in zip(columns, labels):
        vals = {}
        for sector, stats in results.items():
            v = stats.get(col)
            if v is not None:
                vals[sector] = v
        if len(vals) < 2:
            continue
        min_s = min(vals, key=vals.get)
        max_s = max(vals, key=vals.get)
        spread = vals[max_s] - vals[min_s]
        all_v = list(vals.values())
        mean_v = statistics.mean(all_v)
        cv = (statistics.stdev(all_v) / abs(mean_v) * 100) if mean_v != 0 and len(all_v) > 1 else 0
        gaps.append((label, min_s, vals[min_s], max_s, vals[max_s], spread, cv))

    gaps.sort(key=lambda x: -x[6])  # sort by CV%

    for label, min_s, min_v, max_s, max_v, spread, cv in gaps:
        print(f"{label:<25} {min_s:<25} {min_v:>8.2f} {max_s:<25} {max_v:>8.2f} {spread:>8.2f} {cv:>7.1f}%")

    if gaps:
        top = gaps[0]
        print(f"\n>>> MOST VARIABLE COMPONENT: '{top[0]}' with CV={top[6]:.1f}%")
        print(f"    Range: {top[1]} ({top[2]:.2f}) to {top[3]} ({top[4]:.2f}), spread={top[5]:.2f}")


def financial_services_pe_pb_check(data):
    print(f"\n{'=' * 120}")
    print("  FINANCIAL SERVICES PE/PB BIAS CHECK")
    print(f"{'=' * 120}")

    sectors = defaultdict(list)
    for rec in data:
        sectors[rec["sector"]].append(rec)

    print(f"\n{'Sector':<30} {'N':>4} {'Avg PE':>10} {'Avg PB':>10} {'Avg ROE':>10} {'PE_Rel_Sc':>10} {'PB_Rel_Sc':>10} {'Fund_Raw':>10} {'Fund_Final':>10} {'H_Fund':>10}")
    print("-" * 134)

    sector_data = {}
    for sector in sorted(sectors.keys()):
        stocks = sectors[sector]
        n = len(stocks)
        pe_vals = [s["pe_ratio"] for s in stocks if s.get("pe_ratio") is not None and 0 < s["pe_ratio"] < 200]
        pb_vals = [s["pb_ratio"] for s in stocks if s.get("pb_ratio") is not None and 0 < s["pb_ratio"] < 100]
        roe_vals = [s["roe"] for s in stocks if s.get("roe") is not None]
        pe_rel = [s["pe_relative_score"] for s in stocks if s.get("pe_relative_score") is not None]
        pb_rel = [s["pb_relative_score"] for s in stocks if s.get("pb_relative_score") is not None]
        fund_raw = [s["fundamental_score"] for s in stocks if s.get("fundamental_score") is not None]
        fund_final = [s["fundamental_score_final"] for s in stocks if s.get("fundamental_score_final") is not None]
        h_fund = [s["hybrid_fundamental_quality"] for s in stocks if s.get("hybrid_fundamental_quality") is not None]

        def avg(lst):
            return statistics.mean(lst) if lst else None

        row = {
            "pe": avg(pe_vals), "pb": avg(pb_vals), "roe": avg(roe_vals),
            "pe_rel": avg(pe_rel), "pb_rel": avg(pb_rel),
            "fund_raw": avg(fund_raw), "fund_final": avg(fund_final), "h_fund": avg(h_fund),
        }
        sector_data[sector] = row

        def fmt(v):
            return f"{v:>10.2f}" if v is not None else f"{'N/A':>10}"

        print(f"{sector:<30} {n:>4} {fmt(row['pe'])} {fmt(row['pb'])} {fmt(row['roe'])} {fmt(row['pe_rel'])} {fmt(row['pb_rel'])} {fmt(row['fund_raw'])} {fmt(row['fund_final'])} {fmt(row['h_fund'])}")

    fin = sector_data.get("Financial Services")
    if fin:
        print(f"\n--- Financial Services Deep Dive ---")
        print(f"Financial Services avg PE: {fin['pe']:.1f}" if fin['pe'] else "Financial Services avg PE: N/A")
        print(f"Financial Services avg PB: {fin['pb']:.1f}" if fin['pb'] else "Financial Services avg PB: N/A")
        print(f"Financial Services PE relative score: {fin['pe_rel']:.2f}" if fin['pe_rel'] else "Financial Services PE rel: N/A")
        print(f"Financial Services PB relative score: {fin['pb_rel']:.2f}" if fin['pb_rel'] else "Financial Services PB rel: N/A")

        all_pe_rel = [v["pe_rel"] for v in sector_data.values() if v["pe_rel"] is not None]
        all_pb_rel = [v["pb_rel"] for v in sector_data.values() if v["pb_rel"] is not None]
        all_fund = [v["fund_raw"] for v in sector_data.values() if v["fund_raw"] is not None]
        all_fund_final = [v["fund_final"] for v in sector_data.values() if v["fund_final"] is not None]

        if all_pe_rel:
            rank = sorted(all_pe_rel, reverse=True).index(fin["pe_rel"]) + 1 if fin["pe_rel"] is not None else "N/A"
            print(f"PE relative score rank: {rank}/{len(all_pe_rel)} (higher = better score from low PE)")
        if all_pb_rel:
            rank = sorted(all_pb_rel, reverse=True).index(fin["pb_rel"]) + 1 if fin["pb_rel"] is not None else "N/A"
            print(f"PB relative score rank: {rank}/{len(all_pb_rel)} (higher = better score from low PB)")
        if all_fund:
            rank = sorted(all_fund, reverse=True).index(fin["fund_raw"]) + 1 if fin["fund_raw"] is not None else "N/A"
            print(f"Fundamental raw score rank: {rank}/{len(all_fund)}")
        if all_fund_final:
            rank = sorted(all_fund_final, reverse=True).index(fin["fund_final"]) + 1 if fin["fund_final"] is not None else "N/A"
            print(f"Fundamental final score rank: {rank}/{len(all_fund_final)}")

        if fin["pe_rel"] is not None and all_pe_rel:
            mean_pe_rel = statistics.mean(all_pe_rel)
            bias = fin["pe_rel"] - mean_pe_rel
            print(f"\nPE relative score bias vs mean: {bias:+.2f} ({'FAVORS' if bias > 0.5 else 'NEUTRAL' if abs(bias) <= 0.5 else 'DISFAVORS'} Financial Services)")
        if fin["pb_rel"] is not None and all_pb_rel:
            mean_pb_rel = statistics.mean(all_pb_rel)
            bias = fin["pb_rel"] - mean_pb_rel
            print(f"PB relative score bias vs mean: {bias:+.2f} ({'FAVORS' if bias > 0.5 else 'NEUTRAL' if abs(bias) <= 0.5 else 'DISFAVORS'} Financial Services)")


def score_flow_analysis(data):
    """Trace how scores transform through the pipeline per sector."""
    print(f"\n{'=' * 120}")
    print("  SCORE PIPELINE FLOW BY SECTOR")
    print("  (How raw scores transform through adjustments into the final score)")
    print(f"{'=' * 120}")

    pipeline_cols = [
        ("hybrid_overall_score", "H_Overall"),
        ("phase1_blended_score", "Phase1_Blend"),
        ("regime_adjustment_amount", "Regime_Adj"),
        ("sentiment_adjustment_amount", "Sent_Adj"),
        ("volume_adjustment_amount", "Vol_Adj"),
        ("ml_score_adjustment", "ML_Adj"),
        ("crisis_score_adjustment", "Crisis_Adj"),
        ("sector_performance_adj", "Sector_Perf"),
        ("portfolio_context_score", "Port_Ctx"),
        ("final_blended_score", "Final_Blend"),
        ("overall_score", "Overall"),
    ]

    cols = [c[0] for c in pipeline_cols]
    labels = [c[1] for c in pipeline_cols]
    results = sector_stats(data, cols, labels)
    print_table("Score Pipeline", results, cols, labels)

    print(f"\n--- Adjustment Impact Summary ---")
    adj_cols = [
        ("regime_adjustment_amount", "Regime"),
        ("sentiment_adjustment_amount", "Sentiment"),
        ("volume_adjustment_amount", "Volume"),
        ("ml_score_adjustment", "ML"),
        ("crisis_score_adjustment", "Crisis"),
        ("sector_performance_adj", "Sector Perf"),
    ]
    print(f"{'Adjustment':<20} {'Helps Most':<25} {'Value':>8} {'Hurts Most':<25} {'Value':>8} {'Spread':>8}")
    print("-" * 97)

    for col, label in adj_cols:
        vals = {}
        for sector, stats in results.items():
            v = stats.get(col)
            if v is not None:
                vals[sector] = v
        if len(vals) < 2:
            continue
        max_s = max(vals, key=vals.get)
        min_s = min(vals, key=vals.get)
        spread = vals[max_s] - vals[min_s]
        print(f"{label:<20} {max_s:<25} {vals[max_s]:>+8.2f} {min_s:<25} {vals[min_s]:>+8.2f} {spread:>8.2f}")


def hybrid_component_breakdown(data):
    """Break down the hybrid scoring components by sector."""
    print(f"\n{'=' * 120}")
    print("  HYBRID SCORING COMPONENTS BY SECTOR (Core Score Ingredients)")
    print(f"{'=' * 120}")

    hybrid_cols = [
        ("hybrid_fundamental_quality", "H_Fund"),
        ("hybrid_momentum_technical", "H_Momentum"),
        ("hybrid_volume_strength", "H_Volume"),
        ("hybrid_multi_timeframe", "H_MTF"),
        ("hybrid_ml_signal", "H_ML"),
        ("hybrid_risk_adjustment", "H_Risk"),
        ("hybrid_sector_multiplier", "H_SectMult"),
        ("hybrid_overall_score", "H_Overall"),
    ]
    cols = [c[0] for c in hybrid_cols]
    labels = [c[1] for c in hybrid_cols]
    results = sector_stats(data, cols, labels)
    print_table("Hybrid Components", results, cols, labels)
    gap_analysis(results, cols, labels)


def main():
    report = find_latest_report()
    print(f"Analyzing: {os.path.basename(report)}")
    print(f"{'=' * 120}")

    data = load_data(report)
    print(f"Loaded {len(data)} stocks across {len(set(r['sector'] for r in data))} sectors\n")

    # 1. Hybrid component breakdown (core scoring ingredients)
    hybrid_component_breakdown(data)

    # 2. Score pipeline flow (how adjustments change things)
    score_flow_analysis(data)

    # 3. Financial Services PE/PB bias check
    financial_services_pe_pb_check(data)

    # 4. Overall score table with key components
    print(f"\n{'=' * 120}")
    print("  COMPREHENSIVE SECTOR SCORE SUMMARY")
    print(f"{'=' * 120}")

    summary_cols = [
        ("fundamental_score_final", "Fund_Final"),
        ("combined_technical_score_final", "CombTech"),
        ("volume_composite_score", "Vol_Comp"),
        ("institutional_score_final", "Instit"),
        ("portfolio_context_score", "Port_Ctx"),
        ("risk_adjusted_score", "Risk_Adj"),
        ("overall_score", "Overall"),
    ]
    cols = [c[0] for c in summary_cols]
    labels = [c[1] for c in summary_cols]
    results = sector_stats(data, cols, labels)
    print_table("Comprehensive Summary", results, cols, labels)
    gap_analysis(results, cols, labels)

    # 5. Final verdict
    print(f"\n{'=' * 120}")
    print("  FINAL VERDICT: SECTOR BIAS DIAGNOSIS")
    print(f"{'=' * 120}")

    sectors = defaultdict(list)
    for rec in data:
        sectors[rec["sector"]].append(rec)

    overall_by_sector = {}
    for sector, stocks in sectors.items():
        vals = [s["overall_score"] for s in stocks if s.get("overall_score") is not None]
        if vals:
            overall_by_sector[sector] = statistics.mean(vals)

    if overall_by_sector:
        sorted_sectors = sorted(overall_by_sector.items(), key=lambda x: -x[1])
        overall_mean = statistics.mean(overall_by_sector.values())
        print(f"\nOverall score mean across sectors: {overall_mean:.2f}")
        print(f"\nTop 5 scoring sectors:")
        for sector, score in sorted_sectors[:5]:
            print(f"  {sector:<30} {score:.2f} ({score - overall_mean:+.2f} vs mean)")
        print(f"\nBottom 5 scoring sectors:")
        for sector, score in sorted_sectors[-5:]:
            print(f"  {sector:<30} {score:.2f} ({score - overall_mean:+.2f} vs mean)")

        # Identify biggest driver per top/bottom sector
        print(f"\n--- What drives the top sector's advantage? ---")
        if sorted_sectors:
            top_sector = sorted_sectors[0][0]
            components = [
                ("hybrid_fundamental_quality", "Fundamental Quality"),
                ("hybrid_momentum_technical", "Momentum/Technical"),
                ("hybrid_volume_strength", "Volume Strength"),
                ("hybrid_risk_adjustment", "Risk Adjustment"),
                ("regime_adjustment_amount", "Regime Adjustment"),
                ("sentiment_adjustment_amount", "Sentiment Adjustment"),
                ("portfolio_context_score", "Portfolio Context"),
            ]
            for col, label in components:
                top_vals = [s[col] for s in sectors[top_sector] if s.get(col) is not None]
                all_vals = [s[col] for s in data if s.get(col) is not None]
                if top_vals and all_vals:
                    diff = statistics.mean(top_vals) - statistics.mean(all_vals)
                    if abs(diff) > 0.1:
                        direction = "ABOVE" if diff > 0 else "BELOW"
                        print(f"  {label:<25}: {statistics.mean(top_vals):.2f} vs overall {statistics.mean(all_vals):.2f} ({direction} by {abs(diff):.2f})")


if __name__ == "__main__":
    main()
