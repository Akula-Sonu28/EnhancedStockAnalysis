#!/usr/bin/env python3
"""Post-run audit for analyze_top200_stocks_enhanced.py outputs.

Finds the latest report/log, runs contract and safety checks, prints a
markdown summary (PASS/WARN/FAIL). Used by the post-analysis-audit skill
when the user says ANALYSE after a run.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

REPO = Path(__file__).resolve().parents[1]
REPORTS = REPO / "reports"
DATA = REPO / "data"

REQUIRED_SHEETS = (
    "Dashboard",
    "Portfolio Allocation",
    "Past Accuracy",
    "Complete Data",
    "Top Picks",
    "Trading Levels",
    "IC Telemetry",
)

SELL_ACTIONS = frozenset(
    {
        "SELL",
        "WEAK SELL",
        "REDUCE",
        "CONSIDER SELLING",
        "SCALE_OUT_20",
        "SWAP",
        "EXIT",
        "EXIT NOW - Heavy exhaustion",
    }
)

SELL_LIKE_IN_CELL = ("SELL", "WEAK SELL", "REDUCE", "EXIT")


@dataclass
class Check:
    name: str
    status: str  # PASS | WARN | FAIL | INFO
    detail: str


@dataclass
class AuditResult:
    report: Optional[Path] = None
    log: Optional[Path] = None
    merged: Optional[Path] = None
    dashboard: Optional[Path] = None
    checks: List[Check] = field(default_factory=list)
    action_rows: List[Dict[str, Any]] = field(default_factory=list)

    def add(self, name: str, status: str, detail: str) -> None:
        self.checks.append(Check(name, status, detail))


def _latest(glob_pat: str, base: Path) -> Optional[Path]:
    files = sorted(base.glob(glob_pat), key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0] if files else None


def _read_log_tail(path: Path, n: int = 400) -> str:
    try:
        lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError as exc:
        return f"<read error: {exc}>"
    return "\n".join(lines[-n:])


def _ic_metric(ws, label: str) -> Optional[str]:
    for row in ws.iter_rows(max_row=80, values_only=True):
        if row and row[0] == label:
            return str(row[1]) if len(row) > 1 and row[1] is not None else None
    return None


def audit_excel(report: Path, result: AuditResult) -> None:
    try:
        import openpyxl
    except ImportError:
        result.add("Excel audit", "FAIL", "openpyxl not installed")
        return

    try:
        wb = openpyxl.load_workbook(report, read_only=True, data_only=True)
    except Exception as exc:
        result.add("Excel open", "FAIL", str(exc))
        return

    sheets = set(wb.sheetnames)
    missing = [s for s in REQUIRED_SHEETS if s not in sheets]
    if missing:
        result.add("Required sheets", "FAIL", f"missing: {', '.join(missing)}")
    else:
        result.add("Required sheets", "PASS", f"{len(REQUIRED_SHEETS)} contract sheets present")

    if "Top Picks" in sheets:
        ws = wb["Top Picks"]
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        rec_i = next(
            (i for i, h in enumerate(hdr) if h and "recommend" in str(h).lower()),
            None,
        )
        bad = []
        for row in ws.iter_rows(min_row=2, max_row=300, values_only=True):
            if not row or not row[0]:
                continue
            rec = str(row[rec_i] or "") if rec_i is not None else ""
            if any(x in rec.upper() for x in SELL_LIKE_IN_CELL):
                bad.append((row[0], rec))
        if bad:
            result.add("Top Picks no SELL", "FAIL", f"{len(bad)} rows: {bad[:5]}")
        else:
            result.add("Top Picks no SELL", "PASS", "no SELL/REDUCE in sample")

    if "Trading Levels" in sheets:
        ws = wb["Trading Levels"]
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        rec_i = next(
            (i for i, h in enumerate(hdr) if h and "recommend" in str(h).lower()),
            None,
        )
        bad = []
        for row in ws.iter_rows(min_row=2, max_row=300, values_only=True):
            if not row or not row[0]:
                continue
            rec = str(row[rec_i] or "") if rec_i is not None else ""
            if any(x in rec.upper() for x in SELL_LIKE_IN_CELL):
                bad.append((row[0], rec))
        if bad:
            result.add("Trading Levels no SELL", "FAIL", f"{len(bad)} rows")
        else:
            result.add("Trading Levels no SELL", "PASS", "clean")

    if "IC Telemetry" in sheets:
        ws = wb["IC Telemetry"]
        regime = _ic_metric(ws, "current market regime") or _ic_metric(ws, "active regime")
        wf_status = _ic_metric(ws, "walkforward status")
        wf_age = _ic_metric(ws, "walkforward age (days)")
        wf_verdict = _ic_metric(ws, "WALK-FORWARD VERDICT")
        live = _ic_metric(ws, "LIVE ENGINE")
        v2_src = _ic_metric(ws, "v2 weight source")
        ic_gate = _ic_metric(ws, "IC_30d gate (need >= +0.05)")
        if regime:
            result.add("Regime (IC Telemetry)", "INFO", regime)
        if live:
            result.add("Live engine", "INFO", live)
        if v2_src:
            result.add("v2 weights", "INFO", v2_src)
        if wf_status:
            st = "PASS" if wf_status == "FRESH" else "WARN"
            result.add("Walk-forward freshness", st, f"{wf_status} (age {wf_age}d)")
        if wf_verdict:
            st = "PASS" if wf_verdict != "PROMOTE" else "WARN"
            if wf_verdict == "HOLD_SHADOW":
                st = "WARN"
            result.add("Walk-forward verdict", st, wf_verdict)
        if ic_gate:
            result.add("IC_30d gate", "WARN" if ic_gate == "FAIL" else "PASS", ic_gate)

    if "Dashboard" in sheets:
        ws = wb["Dashboard"]
        for row in ws.iter_rows(max_row=5, values_only=True):
            if row and row[0] and "Regime" in str(row[0]):
                result.add("Dashboard header", "INFO", str(row[0])[:120])
                break

    if "Portfolio Allocation" in sheets:
        ws = wb["Portfolio Allocation"]
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        sym_i = next((i for i, h in enumerate(hdr) if h and "WHAT" in str(h).upper()), 0)
        money_i = next((i for i, h in enumerate(hdr) if h and "MONEY" in str(h).upper()), None)
        n_rows = 0
        zero_money = []
        for row in ws.iter_rows(min_row=2, max_row=80, values_only=True):
            if not row or not row[sym_i]:
                continue
            n_rows += 1
            if money_i is not None and row[money_i] in (0, 0.0, "0"):
                zero_money.append(str(row[sym_i]))
        result.add("Portfolio rows", "INFO", str(n_rows))
        if zero_money:
            result.add(
                "Zero MONEY rows",
                "WARN",
                f"{len(zero_money)} symbols (often HOLD=0 deploy): {', '.join(zero_money[:8])}"
                + ("..." if len(zero_money) > 8 else ""),
            )

    wb.close()


def audit_log(log: Path, result: AuditResult) -> None:
    try:
        text = log.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        result.add("Log read", "FAIL", str(exc))
        return

    for pat, label in (
        (r"\bERROR\b", "ERROR lines"),
        (r"\bException\b", "Exception"),
        (r"Traceback", "Traceback"),
    ):
        n = len(re.findall(pat, text))
        if n:
            result.add(label, "FAIL" if label != "ERROR lines" or n > 50 else "WARN", str(n))
        else:
            result.add(label, "PASS", "0")

    if "Batch analysis completed" in text:
        m = re.search(
            r"Batch analysis completed: (\d+) results, (\d+) failures", text
        )
        if m:
            fails = int(m.group(2))
            result.add(
                "Batch completion",
                "PASS" if fails == 0 else "FAIL",
                f"{m.group(1)} results, {fails} failures",
            )
    else:
        result.add("Batch completion", "WARN", "marker not found in log")

    if re.search(r"\[v2\] loaded (\w+) calibrated weights", text):
        m = re.search(
            r"\[v2\] loaded (\w+) calibrated weights from ([^\s]+) \(age=(\d+)d\)",
            text,
        )
        if m:
            result.add(
                "v2 weights loaded",
                "PASS",
                f"{m.group(1)} from {m.group(2)} (age {m.group(3)}d)",
            )
    else:
        result.add("v2 weights loaded", "WARN", "no [v2] loaded line in log")

    cooldown = len(re.findall(r"RECENT_BUY_COOLDOWN", text))
    if cooldown:
        result.add("Buy cooldown suppressions", "INFO", str(cooldown))

    q131_sells = re.findall(
        r"\[Q131-trace\] (\w+): action=(SELL|CONSIDER SELLING|WEAK SELL)[^,]*suppress=(\w+)",
        text,
    )
    active_sells = [s for s, a, sup in q131_sells if sup.lower() == "false"]
    if active_sells:
        result.add("Active exit traces (Q131)", "INFO", ", ".join(active_sells[:12]))


def audit_history(result: AuditResult, run_date: Optional[date] = None) -> None:
    hist = DATA / "recommendation_history.csv"
    if not hist.exists():
        result.add("Recommendation history", "FAIL", "file missing")
        return

    try:
        import pandas as pd
    except ImportError:
        result.add("Recommendation history", "WARN", "pandas required for history checks")
        return

    df = pd.read_csv(hist, low_memory=False)
    date_col = "date" if "date" in df.columns else None
    if not date_col:
        result.add("Recommendation history", "WARN", "no date column")
        return

    df["_dt"] = pd.to_datetime(df[date_col], errors="coerce")
    rd = run_date or date.today()
    today = df[df["_dt"].dt.date == rd]
    result.add("History rows (run date)", "INFO", str(len(today)))

    if today.empty:
        result.add("History today", "WARN", f"no rows for {rd}")
        return

    if "symbol" in today.columns and "action" in today.columns:
        flips = today.groupby("symbol")["action"].nunique()
        multi = flips[flips > 1]
        if len(multi):
            result.add("Same-day action flips", "FAIL", str(dict(multi.head(5))))
        else:
            result.add("Same-day action flips", "PASS", "none")

        sells = today[today["action"].astype(str).str.upper().isin(
            {a.upper() for a in SELL_ACTIONS}
        )]
        for _, row in sells.iterrows():
            result.action_rows.append(
                {
                    "symbol": row.get("symbol"),
                    "action": row.get("action"),
                    "score_v2": row.get("score_v2"),
                    "regime": row.get("regime"),
                }
            )
        if len(sells):
            syms = ", ".join(sells["symbol"].astype(str).tolist())
            result.add("Today sell-side actions", "INFO", syms)


def audit_json_sidecars(result: AuditResult) -> None:
    promo = DATA / "v2_promotion_status.json"
    if promo.exists():
        try:
            p = json.loads(promo.read_text())
            ready = p.get("promotion_ready")
            mode = p.get("mode") or p.get("validation_mode") or "?"
            st = "WARN" if ready and mode == "historical" else "INFO"
            result.add(
                "v2 promotion_status",
                st,
                f"promotion_ready={ready}, mode={mode}",
            )
        except json.JSONDecodeError:
            result.add("v2 promotion_status", "WARN", "invalid JSON")


def audit_dual_strategy_hint(result: AuditResult, terminal_text: Optional[str]) -> None:
    """HIGH CONVICTION is printed to stdout only; parse terminal if provided."""
    sell_syms = {str(r.get("symbol", "")).upper() for r in result.action_rows}
    if not sell_syms:
        return

    if not terminal_text:
        result.add(
            "Dual-strategy HIGH CONVICTION",
            "WARN",
            f"verify terminal: today's sells {', '.join(sorted(sell_syms))} "
            "must NOT appear under HIGH CONVICTION BUY",
        )
        return

    m = re.search(
        r"HIGH CONVICTION \(both strategies say BUY\): (.+)",
        terminal_text,
    )
    if not m:
        result.add("Dual-strategy HIGH CONVICTION", "INFO", "block not in terminal capture")
        return

    conviction = {s.strip().upper() for s in m.group(1).split(",") if s.strip()}
    overlap = sell_syms & conviction
    if overlap:
        result.add(
            "Dual-strategy HIGH CONVICTION",
            "FAIL",
            f"sell symbols listed as BUY conviction: {', '.join(sorted(overlap))}",
        )
    else:
        result.add("Dual-strategy HIGH CONVICTION", "PASS", "no sell/buy contradiction")


def render_markdown(result: AuditResult) -> str:
    lines = [
        "# Post-Analysis Audit",
        "",
        f"**Generated:** {datetime.now().isoformat(timespec='seconds')}",
        "",
        "## Artifacts",
        f"- Report: `{result.report}`" if result.report else "- Report: _not found_",
        f"- Log: `{result.log}`" if result.log else "- Log: _not found_",
        f"- Merged holdings: `{result.merged}`" if result.merged else "",
        f"- Dashboard: `{result.dashboard}`" if result.dashboard else "",
        "",
        "## Checks",
        "",
        "| Check | Status | Detail |",
        "|-------|--------|--------|",
    ]
    for c in result.checks:
        lines.append(f"| {c.name} | **{c.status}** | {c.detail} |")

    if result.action_rows:
        lines.extend(["", "## Today — sell-side (from history)", ""])
        for r in result.action_rows:
            lines.append(
                f"- **{r.get('symbol')}**: {r.get('action')} "
                f"(v2={r.get('score_v2')}, regime={r.get('regime')})"
            )

    fails = [c for c in result.checks if c.status == "FAIL"]
    warns = [c for c in result.checks if c.status == "WARN"]
    lines.extend(
        [
            "",
            "## Summary",
            f"- **FAIL:** {len(fails)}",
            f"- **WARN:** {len(warns)}",
            f"- **PASS/INFO:** {len(result.checks) - len(fails) - len(warns)}",
            "",
        ]
    )
    if fails:
        lines.append("**Investigate FAIL items before acting on recommendations.**")
    elif warns:
        lines.append("**Review WARN items; run may still be usable.**")
    else:
        lines.append("**No blocking failures detected.**")
    return "\n".join(line for line in lines if line is not None)


def run_audit(
    report: Optional[Path] = None,
    log: Optional[Path] = None,
    terminal_text: Optional[str] = None,
    run_date: Optional[date] = None,
) -> AuditResult:
    result = AuditResult()
    result.report = report or _latest("Enhanced_Stock_Report_*.xlsx", REPORTS)
    result.log = log or _latest("top200_analysis_*.log", DATA)
    result.merged = _latest("merged_portfolio_*.xlsx", REPORTS)
    dash = REPO / "Portfolio_Allocation_Dashboard.html"
    result.dashboard = dash if dash.exists() else None

    if not result.report:
        result.add("Latest report", "FAIL", "no Enhanced_Stock_Report_*.xlsx in reports/")
    if not result.log:
        result.add("Latest log", "WARN", "no top200_analysis_*.log in data/")

    if result.report:
        result.add("Latest report file", "INFO", result.report.name)
        audit_excel(result.report, result)
    if result.log:
        result.add("Latest log file", "INFO", result.log.name)
        audit_log(result.log, result)

    audit_history(result, run_date)
    audit_json_sidecars(result)
    audit_dual_strategy_hint(result, terminal_text)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="Post-analysis run audit")
    parser.add_argument("--report", type=Path, help="Excel report path")
    parser.add_argument("--log", type=Path, help="Analysis log path")
    parser.add_argument("--terminal", type=Path, help="Captured terminal stdout")
    parser.add_argument(
        "--out",
        type=Path,
        default=DATA / "post_analysis_audit_latest.md",
        help="Write markdown report here",
    )
    parser.add_argument("--date", help="Run date YYYY-MM-DD (default: today)")
    args = parser.parse_args()

    rd = date.fromisoformat(args.date) if args.date else None
    term = None
    if args.terminal and args.terminal.exists():
        term = args.terminal.read_text(encoding="utf-8", errors="replace")

    result = run_audit(args.report, args.log, term, rd)
    md = render_markdown(result)
    print(md)
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(md, encoding="utf-8")
        print(f"\n_Written to {args.out}_", file=sys.stderr)
    return 1 if any(c.status == "FAIL" for c in result.checks) else 0


if __name__ == "__main__":
    sys.exit(main())
