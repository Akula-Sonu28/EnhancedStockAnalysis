"""Quick fix to correct percentage formatting in existing Excel report"""
import openpyxl
import pandas as pd

# Load the workbook
wb = openpyxl.load_workbook('reports/Enhanced_Stock_Report_20251124_171929.xlsx')
ws = wb['Portfolio Allocation']

# Get headers
headers = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]

# Find percentage columns
percent_columns = ['MY_PROFIT_%', '20D_CHANGE_%', 'PORTFOLIO_%', 'ROE_%', 'VOLATILITY_%', 'BOOK_%_IF_SELL']

print("Fixing percentage columns...")
for col_name in percent_columns:
    if col_name in headers:
        col_idx = headers.index(col_name) + 1  # Excel is 1-indexed
        print(f"\nProcessing {col_name} (column {col_idx})...")
        
        # Fix each data row
        for row_num in range(2, ws.max_row + 1):
            cell = ws.cell(row_num, col_idx)
            value = cell.value
            
            if value is not None and isinstance(value, (int, float)):
                # Always divide by 100 to convert percentage to decimal
                new_value = value / 100
                cell.value = new_value
                cell.number_format = '0.00%'
                
                if row_num <= 5:  # Show first few
                    print(f"  Row {row_num}: {value} → {new_value} (displays as {new_value*100:.2f}%)")

# Save the fixed workbook
output_file = 'reports/Enhanced_Stock_Report_20251124_171929_FIXED.xlsx'
wb.save(output_file)
print(f"\n✅ Fixed report saved to: {output_file}")
print("\nVerifying NATIONALUM row...")

# Verify NATIONALUM
for row_num in range(2, ws.max_row + 1):
    symbol = ws.cell(row_num, 1).value
    if symbol == 'NATIONALUM':
        profit_col = headers.index('MY_PROFIT_%') + 1
        profit_val = ws.cell(row_num, profit_col).value
        print(f"  NATIONALUM MY_PROFIT_%: {profit_val*100:.2f}% (stored as {profit_val})")
        break

print("\n⚠️  Note: Icon set arrows will be corrected in next full report generation.")
print("    The code in analyze_top200_stocks_enhanced.py already has the fix at lines 7313-7322.")
