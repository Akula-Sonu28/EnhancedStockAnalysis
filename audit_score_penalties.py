#!/usr/bin/env python3
"""Audit the latest Excel report for score penalties and unfair adjustments."""

import glob
import os
import numpy as np
import openpyxl

REPORTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "reports")

SEARCH_SUBSTRINGS = [
    "symbol", "sector", "overall_score", "risk_adjusted_score", "final_blended",
    "hybrid_score", "hybrid_overall_score",
    "phase1_adjusted", "phase1_quality", "phase1_context", "phase1_blended",
    "regime_adjusted", "regime_adjustment_amount",
    "sentiment_adjusted", "sentiment_adjustment_amount",
    "sector_performance_adj", "sector_adj_recorded", "quality_adj_recorded",
    "crisis_score_adjustment",
    "portfolio_context_score", "portfolio_fit", "data_quality_score",
    "final_recommendation", "phase2_recommendation",
    "improved_overall_score", "corrected_overall_score",
    "pre_adj_blended_score", "raw_blended_score",
    "volume_adjustment_amount", "volume_adjusted_score",
    "ml_score_adjustment", "sentiment_score_contribution",
    "volume_score_contribution", "pattern_score_contribution",
    "crisis_severity", "crisis_type_detected",
    "score_smoothed", "final_score_with_phase1",
    "hybrid_risk_adjustment", "hybrid_sector_multiplier",
    "fundamental_score", "real_technical_score",
    "optimized_score", "improved_score_used", "hybrid_score_used",
]

ADJUSTMENT_COLS = {
    "regime_adjustment_amount": "Regime Adj",
    "sentiment_adjustment_amount": "Sentiment Adj",
    "sector_performance_adj": "Sector Perf Adj",
    "sector_adj_recorded": "Sector Adj Recorded",
    "quality_adj_recorded": "Quality Adj Recorded",
    "crisis_score_adjustment": "Crisis Adj",
    "phase1_quality_adjustment": "Phase1 Quality Adj",
    "phase1_context_adjustment": "Phase1 Context Adj",
    "volume_adjustment_amount": "Volume Adj",
    "ml_score_adjustment": "ML Adj",
    "sentiment_score_contribution": "Sentiment Contrib",
    "volume_score_contribution": "Volume Contrib",
    "pattern_score_contribution": "Pattern Contrib",
}


def find_latest_report():
    files = sorted(glob.glob(os.path.join(REPORTS_DIR, "*.xlsx")),
                   key=os.path.getmtime, reverse=True)
    if not files:
        raise FileNotFoundError("No .xlsx files found in reports/")
    return files[0]


def match_columns(headers):
    matched = {}
    lower_headers = {i: (h.lower() if h else "") for i, h in enumerate(headers)}
    for idx, lh in lower_headers.items():
        for sub in SEARCH_SUBSTRINGS:
            if sub in lh:
                matched[headers[idx]] = idx
                break
    return matched


def safe_float(val):
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def main():
    report_path = find_latest_report()
    print(f"{'='*100}")
    print(f"SCORE PENALTY & ADJUSTMENT AUDIT")
    print(f"Report: {os.path.basename(report_path)}")
    print(f"{'='*100}\n")

    wb = openpyxl.load_workbook(report_path, read_only=True, data_only=True)
    if "Complete Data" not in wb.sheetnames:
        print("ERROR: 'Complete Data' sheet not found.")
        wb.close()
        return

    ws = wb["Complete Data"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        print("ERROR: No data rows found.")
        return

    headers = list(rows[0])
    col_map = match_columns(headers)

    print(f"Matched {len(col_map)} score-related columns:")
    for name in sorted(col_map.keys()):
        print(f"  [{col_map[name]:3d}] {name}")
    print()

    def get(row, col_name, default=None):
        if col_name in col_map:
            return row[col_map[col_name]]
        return default

    stocks = []
    for row in rows[1:]:
        symbol = get(row, "symbol", "")
        if not symbol:
            continue

        adj_values = {}
        total_penalties = 0.0
        total_bonuses = 0.0
        for col_name in ADJUSTMENT_COLS:
            val = safe_float(get(row, col_name))
            adj_values[col_name] = val
            if val < 0:
                total_penalties += val
            elif val > 0:
                total_bonuses += val

        stock = {
            "symbol": symbol,
            "sector": get(row, "sector", ""),
            "overall_score": safe_float(get(row, "overall_score")),
            "risk_adjusted_score": safe_float(get(row, "risk_adjusted_score")),
            "final_blended_score": safe_float(get(row, "final_blended_score")),
            "hybrid_overall_score": safe_float(get(row, "hybrid_overall_score")),
            "phase1_blended_score": safe_float(get(row, "phase1_blended_score")),
            "pre_adj_blended_score": safe_float(get(row, "pre_adj_blended_score")),
            "raw_blended_score": safe_float(get(row, "raw_blended_score")),
            "improved_overall_score": safe_float(get(row, "improved_overall_score")),
            "corrected_overall_score": safe_float(get(row, "corrected_overall_score")),
            "optimized_score": safe_float(get(row, "optimized_score")),
            "score_smoothed": safe_float(get(row, "score_smoothed")),
            "portfolio_context_score": safe_float(get(row, "portfolio_context_score")),
            "portfolio_fit": get(row, "portfolio_fit", ""),
            "data_quality_score": safe_float(get(row, "data_quality_score")),
            "final_recommendation": get(row, "final_recommendation", ""),
            "phase2_recommendation": get(row, "phase2_recommendation", ""),
            "crisis_type": get(row, "crisis_type_detected", ""),
            "crisis_severity": safe_float(get(row, "crisis_severity")),
            "fundamental_score": safe_float(get(row, "fundamental_score")),
            "real_technical_score": safe_float(get(row, "real_technical_score")),
            "adjustments": adj_values,
            "total_penalties": total_penalties,
            "total_bonuses": total_bonuses,
            "net_adjustment": total_bonuses + total_penalties,
        }
        stocks.append(stock)

    print(f"Total stocks analyzed: {len(stocks)}\n")

    # ── 1. TOP 15 by final_blended_score ──
    print(f"{'='*100}")
    print("TOP 15 STOCKS BY FINAL BLENDED SCORE")
    print(f"{'='*100}")
    by_final = sorted(stocks, key=lambda s: s["final_blended_score"], reverse=True)[:15]
    print(f"{'Symbol':<18} {'Sector':<22} {'Hybrid':>7} {'PreAdj':>7} {'Final':>7} "
          f"{'Penalties':>10} {'Bonuses':>8} {'Net':>7} {'Recommendation':<18}")
    print("-" * 130)
    for s in by_final:
        print(f"{str(s['symbol']):<18} {str(s['sector'])[:21]:<22} "
              f"{s['hybrid_overall_score']:7.2f} {s['pre_adj_blended_score']:7.2f} "
              f"{s['final_blended_score']:7.2f} "
              f"{s['total_penalties']:10.2f} {s['total_bonuses']:8.2f} "
              f"{s['net_adjustment']:7.2f} {str(s['final_recommendation']):<18}")
    print()

    for s in by_final[:5]:
        print(f"  {s['symbol']} adjustment breakdown:")
        for col_name, label in ADJUSTMENT_COLS.items():
            val = s["adjustments"][col_name]
            if val != 0:
                marker = "⊖" if val < 0 else "⊕"
                print(f"    {marker} {label:<25} {val:+.4f}")
        print()

    # ── 2. TOP 15 MOST PENALIZED ──
    print(f"{'='*100}")
    print("TOP 15 MOST-PENALIZED STOCKS (largest negative total)")
    print(f"{'='*100}")
    by_penalty = sorted(stocks, key=lambda s: s["total_penalties"])[:15]
    print(f"{'Symbol':<18} {'Sector':<22} {'Hybrid':>7} {'Final':>7} "
          f"{'Penalties':>10} {'Regime':>8} {'Sent':>8} {'SectAdj':>8} "
          f"{'QualAdj':>8} {'Crisis':>8} {'VolAdj':>8}")
    print("-" * 140)
    for s in by_penalty:
        a = s["adjustments"]
        print(f"{str(s['symbol']):<18} {str(s['sector'])[:21]:<22} "
              f"{s['hybrid_overall_score']:7.2f} {s['final_blended_score']:7.2f} "
              f"{s['total_penalties']:10.2f} "
              f"{a.get('regime_adjustment_amount',0):8.4f} "
              f"{a.get('sentiment_adjustment_amount',0):8.4f} "
              f"{a.get('sector_adj_recorded',0):8.4f} "
              f"{a.get('quality_adj_recorded',0):8.4f} "
              f"{a.get('crisis_score_adjustment',0):8.4f} "
              f"{a.get('volume_adjustment_amount',0):8.4f}")
    print()

    # ── 3. CORRELATION: hybrid_overall_score vs final_blended_score ──
    print(f"{'='*100}")
    print("CORRELATION ANALYSIS")
    print(f"{'='*100}")
    hybrid_vals = np.array([s["hybrid_overall_score"] for s in stocks if s["hybrid_overall_score"] > 0])
    final_vals = np.array([s["final_blended_score"] for s in stocks if s["hybrid_overall_score"] > 0])

    if len(hybrid_vals) > 2:
        corr = np.corrcoef(hybrid_vals, final_vals)[0, 1]
        print(f"  Pearson correlation (hybrid_overall_score vs final_blended_score): {corr:.4f}")
        if corr < 0.90:
            print(f"  ⚠️  ALERT: Correlation {corr:.4f} < 0.90 — adjustments are DISTORTING the core signal!")
        else:
            print(f"  ✓  Correlation {corr:.4f} >= 0.90 — adjustments preserve core signal ranking.")
    print()

    pre_adj_vals = np.array([s["pre_adj_blended_score"] for s in stocks if s["pre_adj_blended_score"] > 0])
    final_pre = np.array([s["final_blended_score"] for s in stocks if s["pre_adj_blended_score"] > 0])
    if len(pre_adj_vals) > 2:
        corr2 = np.corrcoef(pre_adj_vals, final_pre)[0, 1]
        print(f"  Pearson correlation (pre_adj_blended vs final_blended_score):     {corr2:.4f}")
    print()

    # ── 4. Good stocks killed by penalties ──
    print(f"{'='*100}")
    print("GOOD STOCKS KILLED BY PENALTIES (hybrid > 60, final < 50)")
    print(f"{'='*100}")
    killed = [s for s in stocks if s["hybrid_overall_score"] > 60 and s["final_blended_score"] < 50]
    if killed:
        print(f"Found {len(killed)} stocks unfairly penalized:")
        print(f"{'Symbol':<18} {'Hybrid':>7} {'Final':>7} {'Penalties':>10} {'Net':>7} {'Recommendation':<18}")
        print("-" * 80)
        for s in sorted(killed, key=lambda x: x["total_penalties"]):
            print(f"{str(s['symbol']):<18} {s['hybrid_overall_score']:7.2f} "
                  f"{s['final_blended_score']:7.2f} {s['total_penalties']:10.2f} "
                  f"{s['net_adjustment']:7.2f} {str(s['final_recommendation']):<18}")
            for col_name, label in ADJUSTMENT_COLS.items():
                val = s["adjustments"][col_name]
                if val < 0:
                    print(f"    ⊖ {label:<25} {val:+.4f}")
    else:
        print("  None found — no good stocks being killed by penalties.")
    print()

    # ── 5. Bad stocks inflated ──
    print(f"{'='*100}")
    print("BAD STOCKS INFLATED (hybrid < 40, final > 50)")
    print(f"{'='*100}")
    inflated = [s for s in stocks if s["hybrid_overall_score"] < 40 and s["final_blended_score"] > 50]
    if inflated:
        print(f"Found {len(inflated)} stocks unfairly inflated:")
        print(f"{'Symbol':<18} {'Hybrid':>7} {'Final':>7} {'Bonuses':>8} {'Net':>7} {'Recommendation':<18}")
        print("-" * 80)
        for s in sorted(inflated, key=lambda x: x["total_bonuses"], reverse=True):
            print(f"{str(s['symbol']):<18} {s['hybrid_overall_score']:7.2f} "
                  f"{s['final_blended_score']:7.2f} {s['total_bonuses']:8.2f} "
                  f"{s['net_adjustment']:7.2f} {str(s['final_recommendation']):<18}")
            for col_name, label in ADJUSTMENT_COLS.items():
                val = s["adjustments"][col_name]
                if val > 0:
                    print(f"    ⊕ {label:<25} {val:+.4f}")
    else:
        print("  None found — no bad stocks being inflated.")
    print()

    # ── 6. Average adjustments by type ──
    print(f"{'='*100}")
    print("AVERAGE ADJUSTMENT BY TYPE (across all stocks)")
    print(f"{'='*100}")
    n = len(stocks)
    for col_name, label in ADJUSTMENT_COLS.items():
        vals = [s["adjustments"][col_name] for s in stocks]
        avg = sum(vals) / n if n else 0
        mn = min(vals) if vals else 0
        mx = max(vals) if vals else 0
        neg_count = sum(1 for v in vals if v < 0)
        pos_count = sum(1 for v in vals if v > 0)
        zero_count = sum(1 for v in vals if v == 0)
        print(f"  {label:<25} avg={avg:+8.4f}  min={mn:+8.4f}  max={mx:+8.4f}  "
              f"neg={neg_count:3d}  pos={pos_count:3d}  zero={zero_count:3d}")
    print()

    # ── 7. Score journey summary ──
    print(f"{'='*100}")
    print("SCORE JOURNEY SUMMARY (avg across all stocks)")
    print(f"{'='*100}")
    stages = [
        ("hybrid_overall_score", "Hybrid Overall Score"),
        ("improved_overall_score", "Improved Overall Score"),
        ("corrected_overall_score", "Corrected Overall Score"),
        ("pre_adj_blended_score", "Pre-Adjustment Blended"),
        ("raw_blended_score", "Raw Blended Score"),
        ("score_smoothed", "Smoothed Score"),
        ("phase1_blended_score", "Phase1 Blended Score"),
        ("final_blended_score", "Final Blended Score"),
        ("overall_score", "Overall Score (final)"),
        ("risk_adjusted_score", "Risk-Adjusted Score"),
    ]
    for key, label in stages:
        vals = [s[key] for s in stocks if s[key] > 0]
        if vals:
            print(f"  {label:<30} avg={np.mean(vals):7.2f}  min={min(vals):7.2f}  "
                  f"max={max(vals):7.2f}  std={np.std(vals):6.2f}")
    print()

    # ── 8. Distribution of net adjustments ──
    print(f"{'='*100}")
    print("NET ADJUSTMENT DISTRIBUTION")
    print(f"{'='*100}")
    nets = [s["net_adjustment"] for s in stocks]
    bins = [(-999, -10), (-10, -5), (-5, -2), (-2, 0), (0, 0.001), (0.001, 2), (2, 5), (5, 10), (10, 999)]
    labels = ["< -10", "-10 to -5", "-5 to -2", "-2 to 0", "exactly 0", "0 to +2", "+2 to +5", "+5 to +10", "> +10"]
    for (lo, hi), lbl in zip(bins, labels):
        cnt = sum(1 for v in nets if lo <= v < hi)
        bar = "█" * cnt
        print(f"  {lbl:>12}: {cnt:3d}  {bar}")
    print()

    # ── 9. Recommendation distribution ──
    print(f"{'='*100}")
    print("FINAL RECOMMENDATION DISTRIBUTION")
    print(f"{'='*100}")
    rec_counts = {}
    for s in stocks:
        rec = str(s["final_recommendation"])
        rec_counts[rec] = rec_counts.get(rec, 0) + 1
    for rec, cnt in sorted(rec_counts.items(), key=lambda x: -x[1]):
        avg_final = np.mean([s["final_blended_score"] for s in stocks
                             if str(s["final_recommendation"]) == rec])
        avg_hybrid = np.mean([s["hybrid_overall_score"] for s in stocks
                              if str(s["final_recommendation"]) == rec])
        print(f"  {rec:<25} count={cnt:3d}  avg_hybrid={avg_hybrid:6.2f}  avg_final={avg_final:6.2f}")
    print()

    print(f"{'='*100}")
    print("AUDIT COMPLETE")
    print(f"{'='*100}")


if __name__ == "__main__":
    main()
